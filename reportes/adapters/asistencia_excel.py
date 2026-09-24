from core.utils.asistencia import calcular_estado_asistencia
from datetime import datetime, timedelta
from asistencia.models import Asistencia, Movimiento
from core.models import IncidenciaDia
from core.calculadora import CalculadoraAsistencia
from openpyxl import Workbook
from copy import copy
from openpyxl.utils import get_column_letter

from core.excel.encabezado import escribir_encabezado_reporte
from core.excel.impresion import configurar_impresion
from core.excel.respuesta import crear_respuesta_excel
from core.excel.estilos import (
    ALINEACION_CENTRO,
    ALINEACION_IZQUIERDA,
    BORDE_FINO,
    FUENTE_ENCABEZADO,
    FUENTE_NEGRITA,
    RELLENO_ALERTA,
    RELLENO_ERROR,
    RELLENO_GRIS,
    RELLENO_TITULO,
)

from reportes.services.asistencia import (
    aplicar_filtros_asistencia,
    calcular_incidencias_asistencia,
    obtener_asistencias_base,
)
from core.excel.estilos import (
    RELLENO_ALERTA,
    RELLENO_ERROR,
    RELLENO_GRIS,
    RELLENO_INFORMATIVO,
    RELLENO_OK,
)
from core.utils.asistencia import calcular_estado_asistencia
from reportes.services.asistencia import (
    aplicar_filtros_asistencia,
    calcular_incidencias_asistencia,
    obtener_asistencias_base,
)

COLORES_ESTADO = {
    "OK": RELLENO_OK,
    "ASISTENCIA": RELLENO_OK,
    "COMPLETO": RELLENO_OK,
    "RETARDO": RELLENO_ALERTA,
    "FALTA": RELLENO_ERROR,
    "INCOMPLETO": RELLENO_INFORMATIVO,
    "PENDIENTE": RELLENO_INFORMATIVO,
    "TIEMPO_EXTRA": RELLENO_INFORMATIVO,
    "VACACIONES": RELLENO_GRIS,
    "INCAPACIDAD": RELLENO_GRIS,
    "DESCANSO": RELLENO_GRIS,
    "PERMISO": RELLENO_GRIS,
    "INCIDENCIA": RELLENO_GRIS,
    "NO_LABORAL": RELLENO_GRIS,
    "SIN CONTROL": RELLENO_GRIS,
    "SIN_TURNO": RELLENO_GRIS,
}


COLORES_INCIDENCIA = {
    "OK": RELLENO_OK,
    "SIN SALIDA": RELLENO_ERROR,
    "RETARDO": RELLENO_ALERTA,
    "SIN TURNO": RELLENO_GRIS,
}


def construir_reporte_asistencia(request):
    empresa = request.empresa
    inicio = request.GET.get("inicio")
    fin = request.GET.get("fin")

    registros = obtener_asistencias_base(request)
    registros = aplicar_filtros_asistencia(
        request,
        registros,
    ).order_by(
        "empleado__numero_empleado",
        "-fecha",
    )

    filas = []

    for registro in registros:
        estado = calcular_estado_asistencia(
            registro.empleado,
            registro.fecha,
        )

        incidencias = calcular_incidencias_asistencia(
            registro.empleado,
            registro.fecha,
        )

        incidencia_texto = " | ".join(incidencias) if incidencias else "OK"

        filas.append(
            [
                registro.empleado.numero_empleado,
                registro.empleado.nombre,
                registro.fecha,
                registro.hora_entrada or "--",
                registro.hora_salida or "--",
                estado,
                incidencia_texto,
            ]
        )

    return {
        "titulo": "REPORTE DE ASISTENCIA",
        "empresa": empresa,
        "nombre_hoja": "Asistencia",
        "nombre_archivo": "reporte_asistencia.xlsx",
        "inicio": inicio,
        "fin": fin,
        "encabezados": [
            "No. Empleado",
            "Empleado",
            "Fecha",
            "Hora Entrada",
            "Hora Salida",
            "Estado",
            "Incidencias",
        ],
        "filas": filas,
        "anchos": {
            1: 14,
            2: 35,
            3: 13,
            4: 14,
            5: 14,
            6: 16,
            7: 38,
        },
        "formatos": {
            3: "dd/mm/yyyy",
            4: "hh:mm",
            5: "hh:mm",
        },
        "columnas_centradas": [
            1,
            3,
            4,
            5,
            6,
        ],
        "columnas_envueltas": [
            7,
        ],
        "reglas_color": {
            6: COLORES_ESTADO,
            7: COLORES_INCIDENCIA,
        },
    }


def crear_reporte_asistencia_semanal(request):
    """
    Genera los datos del reporte semanal de asistencia.

    Semana operativa:
    jueves a miércoles.
    """

    from nucleo.models import Departamento, Empleado

    empresa = request.empresa
    inicio = request.GET.get("inicio")

    # ==========================================
    # SEMANA OPERATIVA: JUEVES A MIERCOLES
    # ==========================================
    dias_semana = []

    if inicio:
        fecha_seleccionada = datetime.strptime(
            inicio,
            "%Y-%m-%d",
        ).date()

        dias_desde_jueves = (fecha_seleccionada.weekday() - 3) % 7

        fecha_inicio = fecha_seleccionada - timedelta(days=dias_desde_jueves)

        dias_semana = [fecha_inicio + timedelta(days=i) for i in range(7)]

    # ==========================================
    # EMPLEADOS
    # ==========================================
    empleados = Empleado.objects.filter(
        empresa=empresa,
        activo=True,
    )

    departamento_id = request.GET.get("departamento")
    empleado_id = request.GET.get("empleado")

    departamentos = Departamento.objects.filter(
        empresa=empresa,
        activo=True,
    ).order_by("nombre")

    if departamento_id and departamento_id not in ["", "0"]:
        departamentos = departamentos.filter(id=departamento_id)

    if departamento_id and departamento_id not in ["", "0"]:
        empleados = empleados.filter(departamento_id=departamento_id)

    if empleado_id and empleado_id not in ["", "0"]:
        empleados = empleados.filter(id=empleado_id)

    empleados = empleados.order_by("numero_empleado")

    # ==========================================
    # ASISTENCIAS
    # ==========================================
    asistencias_semana = {}

    if dias_semana:
        asistencias_qs = Asistencia.objects.filter(
            empresa=empresa,
            empleado__in=empleados,
            fecha__range=(
                dias_semana[0],
                dias_semana[-1],
            ),
        )

        asistencias_semana = {
            (
                asistencia.empleado_id,
                asistencia.fecha,
            ): asistencia
            for asistencia in asistencias_qs
        }

    # ==========================================
    # MOVIMIENTOS
    # ==========================================
    movimientos_semana = {}

    if dias_semana:
        movimientos_qs = Movimiento.objects.filter(
            asistencia__empresa=empresa,
            asistencia__empleado__in=empleados,
            fecha__range=(
                dias_semana[0],
                dias_semana[-1],
            ),
        ).select_related(
            "asistencia",
            "asistencia__empleado",
        )

        for movimiento in movimientos_qs:
            clave = (
                movimiento.asistencia.empleado_id,
                movimiento.fecha,
            )

            movimientos_semana.setdefault(
                clave,
                [],
            ).append(movimiento)

    # ==========================================
    # INCIDENCIAS
    # ==========================================
    incidencias_semana = {}

    if dias_semana:
        incidencias_qs = IncidenciaDia.objects.filter(
            empleado__in=empleados,
            fecha__range=(
                dias_semana[0],
                dias_semana[-1],
            ),
        ).select_related(
            "empleado",
            "incidencia",
        )

        incidencias_semana = {
            (
                incidencia.empleado_id,
                incidencia.fecha,
            ): incidencia
            for incidencia in incidencias_qs
        }

    # ==========================================
    # MATRIZ SEMANAL
    # ==========================================
    matriz_semanal = []

    for empleado_matriz in empleados:
        fila = {
            "empleado": empleado_matriz,
            "dias": [],
        }

        for fecha_dia in dias_semana:
            clave = (
                empleado_matriz.id,
                fecha_dia,
            )

            asistencia_dia = asistencias_semana.get(clave)
            movimientos_dia = movimientos_semana.get(
                clave,
                [],
            )
            incidencia_dia = incidencias_semana.get(clave)

            calculadora = CalculadoraAsistencia(
                empleado_matriz,
                fecha_dia,
                movimientos=movimientos_dia,
                asistencia=asistencia_dia,
                incidencia_dia=incidencia_dia,
            )

            resultado = calculadora.calcular()

            fila["dias"].append(
                {
                    "fecha": fecha_dia,
                    "entrada": resultado.get("entrada"),
                    "salida": resultado.get("salida"),
                    "estado": resultado.get("estado"),
                    "tipo_incidencia": resultado.get("tipo_incidencia"),
                    "es_retardo": resultado.get("es_retardo"),
                    "es_falta": resultado.get("es_falta"),
                    "es_incompleto": resultado.get("es_incompleto"),
                    "es_irregular": resultado.get("es_irregular"),
                }
            )

        matriz_semanal.append(fila)

    # ==========================================
    # CREAR LIBRO Y HOJA
    # ==========================================
    wb = Workbook()
    hoja_inicial = wb.active
    wb.remove(hoja_inicial)

    for departamento_reporte in departamentos:

        if departamento_reporte:
            ws = wb.create_sheet(title=departamento_reporte.nombre[:31])
        else:
            ws = wb.create_sheet(title="Asistencia")

        escribir_encabezado_reporte(
            ws=ws,
            titulo="REPORTE SEMANAL DE ASISTENCIA",
            empresa=empresa,
            inicio=dias_semana[0] if dias_semana else None,
            fin=dias_semana[-1] if dias_semana else None,
            ultima_columna="P",
        )

        if departamento_reporte:
            ws.merge_cells("A5:P5")
            celda_departamento = ws["A5"]
            celda_departamento.value = f"Departamento: {departamento_reporte.nombre}"
            celda_departamento.font = FUENTE_NEGRITA
            celda_departamento.alignment = ALINEACION_IZQUIERDA

        # ==========================================
        # ENCABEZADOS DE DOS NIVELES
        # ==========================================
        fila_dias = 6
        fila_horas = 7

        ws.merge_cells(
            start_row=fila_dias,
            start_column=1,
            end_row=fila_horas,
            end_column=1,
        )
        ws.cell(
            row=fila_dias,
            column=1,
            value="No.",
        )

        ws.merge_cells(
            start_row=fila_dias,
            start_column=2,
            end_row=fila_horas,
            end_column=2,
        )
        ws.cell(
            row=fila_dias,
            column=2,
            value="Empleado",
        )

        nombres_dias = [
            "Jueves",
            "Viernes",
            "Sabado",
            "Domingo",
            "Lunes",
            "Martes",
            "Miercoles",
        ]

        columna = 3

        for nombre_dia, fecha_dia in zip(
            nombres_dias,
            dias_semana,
        ):
            ws.merge_cells(
                start_row=fila_dias,
                start_column=columna,
                end_row=fila_dias,
                end_column=columna + 1,
            )

            ws.cell(
                row=fila_dias,
                column=columna,
                value=(f"{nombre_dia}\n" f"{fecha_dia.strftime('%d/%m/%Y')}"),
            )

            ws.cell(
                row=fila_horas,
                column=columna,
                value="Entrada",
            )

            ws.cell(
                row=fila_horas,
                column=columna + 1,
                value="Salida",
            )

            columna += 2

        # Aplicar estilo a las 16 columnas
        for fila in range(fila_dias, fila_horas + 1):
            for columna in range(1, 17):
                celda = ws.cell(
                    row=fila,
                    column=columna,
                )
                celda.font = FUENTE_ENCABEZADO
                celda.fill = RELLENO_TITULO
                celda.alignment = ALINEACION_CENTRO
                celda.border = BORDE_FINO

        ws.row_dimensions[fila_dias].height = 34
        ws.row_dimensions[fila_horas].height = 22

        # Anchos iniciales
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 30

        for columna in range(3, 17):
            ws.column_dimensions[get_column_letter(columna)].width = 11

        # ==========================================
        # DATOS DE EMPLEADOS
        # ==========================================
        fila_actual = 8

        matriz_departamento = [
            fila
            for fila in matriz_semanal
            if (
                departamento_reporte
                and fila["empleado"].departamento_id == departamento_reporte.id
            )
        ]

        for fila_matriz in matriz_departamento:
            empleado = fila_matriz["empleado"]

            celda_numero = ws.cell(
                row=fila_actual,
                column=1,
                value=empleado.numero_empleado,
            )
            celda_numero.alignment = ALINEACION_CENTRO
            celda_numero.border = BORDE_FINO

            celda_nombre = ws.cell(
                row=fila_actual,
                column=2,
                value=empleado.nombre,
            )
            celda_nombre.alignment = ALINEACION_IZQUIERDA
            celda_nombre.border = BORDE_FINO

            columna = 3

            for dia in fila_matriz["dias"]:

                # Incidencia: ocupa Entrada + Salida
                if dia["tipo_incidencia"]:
                    ws.merge_cells(
                        start_row=fila_actual,
                        start_column=columna,
                        end_row=fila_actual,
                        end_column=columna + 1,
                    )

                    celda = ws.cell(
                        row=fila_actual,
                        column=columna,
                        value=dia["tipo_incidencia"],
                    )
                    celda.fill = RELLENO_GRIS
                    celda.font = FUENTE_NEGRITA
                    celda.alignment = ALINEACION_CENTRO

                # Falta: ocupa Entrada + Salida
                elif dia["es_falta"]:
                    ws.merge_cells(
                        start_row=fila_actual,
                        start_column=columna,
                        end_row=fila_actual,
                        end_column=columna + 1,
                    )

                    celda = ws.cell(
                        row=fila_actual,
                        column=columna,
                        value="FALTA",
                    )
                    celda.fill = RELLENO_ERROR
                    celda.font = FUENTE_NEGRITA
                    celda.alignment = ALINEACION_CENTRO

                # Día no laboral
                elif dia["estado"] == "NO_LABORAL":
                    ws.merge_cells(
                        start_row=fila_actual,
                        start_column=columna,
                        end_row=fila_actual,
                        end_column=columna + 1,
                    )

                    celda = ws.cell(
                        row=fila_actual,
                        column=columna,
                        value="NO LABORAL",
                    )
                    celda.fill = RELLENO_GRIS
                    celda.alignment = ALINEACION_CENTRO

                # Día futuro
                elif dia["estado"] == "FUTURO":
                    ws.merge_cells(
                        start_row=fila_actual,
                        start_column=columna,
                        end_row=fila_actual,
                        end_column=columna + 1,
                    )

                    celda = ws.cell(
                        row=fila_actual,
                        column=columna,
                        value="--",
                    )
                    celda.alignment = ALINEACION_CENTRO

                # Día con registros
                else:
                    entrada = ws.cell(
                        row=fila_actual,
                        column=columna,
                        value=dia["entrada"] or "--",
                    )

                    salida = ws.cell(
                        row=fila_actual,
                        column=columna + 1,
                        value=dia["salida"] or "--",
                    )

                    entrada.alignment = ALINEACION_CENTRO
                    salida.alignment = ALINEACION_CENTRO

                    entrada.border = BORDE_FINO
                    salida.border = BORDE_FINO

                    if dia["entrada"]:
                        entrada.number_format = "hh:mm"

                    if dia["salida"]:
                        salida.number_format = "hh:mm"

                    if dia["es_retardo"]:
                        entrada.fill = RELLENO_ALERTA
                        entrada.font = FUENTE_NEGRITA

                    if dia["es_incompleto"] or dia["es_irregular"]:
                        if not dia["entrada"]:
                            entrada.value = "FALTA REG."

                        if not dia["salida"]:
                            salida.value = "FALTA REG."

                # Asegurar bordes en ambas posiciones del día
                for numero_columna in (
                    columna,
                    columna + 1,
                ):
                    ws.cell(
                        row=fila_actual,
                        column=numero_columna,
                    ).border = BORDE_FINO

                columna += 2

            fila_actual += 1

    return crear_respuesta_excel(
        workbook=wb,
        nombre_archivo="reporte_semanal_asistencia.xlsx",
    )
